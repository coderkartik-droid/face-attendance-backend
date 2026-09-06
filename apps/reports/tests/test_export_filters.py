"""Export endpoints must query the DB fresh and honour report filters."""

from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import StudentProfile, TeacherProfile, User
from apps.academics.models import Class, Section
from apps.attendance.models import AttendanceRecord, AttendanceSession


class ExportFilterTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="admin",
            password="pass12345",
            role="school_admin",
            first_name="Admin",
            last_name="User",
        )
        self.cls = Class.objects.create(name="Grade 10", code="G10")
        self.sec = Section.objects.create(name="A", class_obj=self.cls)
        self.teacher = User.objects.create_user(
            username="t1",
            password="pass12345",
            role="teacher",
            first_name="Teach",
            last_name="Er",
        )
        TeacherProfile.objects.create(user=self.teacher, employee_id="T1")

        today = timezone.localdate()
        yesterday = today - timedelta(days=1)
        self.session_today, _ = AttendanceSession.objects.get_or_create(
            class_obj=self.cls,
            section_obj=self.sec,
            date=today,
            session_name="Daily Attendance",
            defaults={"teacher": self.teacher},
        )
        self.session_yesterday, _ = AttendanceSession.objects.get_or_create(
            class_obj=self.cls,
            section_obj=self.sec,
            date=yesterday,
            session_name="Daily Attendance",
            defaults={"teacher": self.teacher},
        )

        self.alice = User.objects.create_user(
            username="alice",
            password="pass12345",
            role="student",
            first_name="Alice",
            last_name="A",
        )
        StudentProfile.objects.create(
            user=self.alice,
            roll_number="R1",
            class_obj=self.cls,
            section_obj=self.sec,
        )
        self.bob = User.objects.create_user(
            username="bob",
            password="pass12345",
            role="student",
            first_name="Bob",
            last_name="B",
        )
        StudentProfile.objects.create(
            user=self.bob,
            roll_number="R2",
            class_obj=self.cls,
            section_obj=self.sec,
        )

        self._mark(self.session_today, self.alice, AttendanceRecord.Status.PRESENT)
        self._mark(self.session_today, self.bob, AttendanceRecord.Status.ABSENT)
        self._mark(self.session_yesterday, self.alice, AttendanceRecord.Status.LATE)

        self.client.force_authenticate(self.admin)

    def _xlsx_text(self, response):
        wb = load_workbook(BytesIO(response.content))
        rows = []
        for row in wb.active.iter_rows(values_only=True):
            rows.extend("" if cell is None else str(cell) for cell in row)
        return "\n".join(rows)

    def _snapshot(self, user, session):
        profile = getattr(user, "student_profile", None) or getattr(
            user, "teacher_profile", None
        )
        return {
            "student_name": user.full_name,
            "student_role": user.role,
            "roll_number": getattr(profile, "roll_number", "") or "",
            "employee_id": getattr(profile, "employee_id", "") or "",
            "class_name": session.class_obj.name,
            "section_name": session.section_obj.name,
        }

    def _mark(self, session, user, status):
        AttendanceRecord.objects.update_or_create(
            session=session,
            student=user,
            defaults={
                "status": status,
                "verification_method": AttendanceRecord.Method.MANUAL,
                **self._snapshot(user, session),
            },
        )

    def test_default_export_is_today_only(self):
        response = self.client.get("/api/reports/export/excel/")
        self.assertEqual(response.status_code, 200)
        body = self._xlsx_text(response)
        self.assertIn("Alice", body)
        self.assertIn("Bob", body)
        # Yesterday's LATE row for Alice must not appear in an unfiltered export.
        self.assertNotIn("LATE", body)

    def test_custom_range_excludes_outside_dates(self):
        yesterday = timezone.localdate() - timedelta(days=1)
        response = self.client.get(
            "/api/reports/export/excel/",
            {"start_date": str(yesterday), "end_date": str(yesterday)},
        )
        self.assertEqual(response.status_code, 200)
        body = self._xlsx_text(response)
        self.assertIn("Alice", body)
        self.assertIn("LATE", body)
        self.assertNotIn("Bob", body)

    def test_status_and_search_filters(self):
        today = timezone.localdate()
        response = self.client.get(
            "/api/reports/export/pdf/",
            {"date": str(today), "status": "PRESENT", "search": "Alice"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertGreater(len(response.content), 100)

    def test_second_export_picks_up_new_row(self):
        """A later export must not reuse the first request's queryset."""
        carol = User.objects.create_user(
            username="carol",
            password="pass12345",
            role="student",
            first_name="Carol",
            last_name="C",
        )
        StudentProfile.objects.create(
            user=carol,
            roll_number="R3",
            class_obj=self.cls,
            section_obj=self.sec,
        )
        first = self.client.get("/api/reports/export/excel/")
        self.assertNotIn("Carol", self._xlsx_text(first))

        self._mark(self.session_today, carol, AttendanceRecord.Status.PRESENT)
        second = self.client.get("/api/reports/export/excel/")
        self.assertIn("Carol", self._xlsx_text(second))
